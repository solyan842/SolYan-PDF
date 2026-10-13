from __future__ import annotations

import io
import json
import os
import shutil
import tempfile
import threading
import traceback
from copy import copy
from pathlib import Path
from typing import Any

from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.responses import Response

MAX_FILE_SIZE = int(os.getenv("OCR_MAX_FILE_SIZE", str(50 * 1024 * 1024)))
SUPPORTED_SUFFIXES = {".pdf", ".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp"}
IMAGE_SUFFIXES = SUPPORTED_SUFFIXES - {".pdf"}
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PRINT_LONG_EDGE_INCHES = 14.0
TABLE_GAP_ROWS = 2

app = FastAPI(title="SolYan OCR Service", version="0.3.0")

_ocr_pipeline: Any | None = None
_structure_pipeline: Any | None = None
_model_lock = threading.Lock()


def _load_ocr() -> Any:
    global _ocr_pipeline
    if _ocr_pipeline is None:
        with _model_lock:
            if _ocr_pipeline is None:
                from paddleocr import PaddleOCR

                _ocr_pipeline = PaddleOCR(
                    lang=os.getenv("OCR_LANG", "vi"),
                    use_doc_orientation_classify=True,
                    use_doc_unwarping=True,
                    use_textline_orientation=True,
                    engine="paddle",
                )
    return _ocr_pipeline


def _load_structure() -> Any:
    global _structure_pipeline
    if _structure_pipeline is None:
        with _model_lock:
            if _structure_pipeline is None:
                from paddleocr import PPStructureV3

                _structure_pipeline = PPStructureV3(
                    use_doc_orientation_classify=True,
                    use_doc_unwarping=True,
                    engine="paddle",
                )
    return _structure_pipeline


def _read_upload(upload: UploadFile) -> tuple[bytes, str]:
    suffix = Path(upload.filename or "document.bin").suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise HTTPException(status_code=400, detail=f"Unsupported input type: {suffix or 'unknown'}")
    data = upload.file.read(MAX_FILE_SIZE + 1)
    if len(data) > MAX_FILE_SIZE:
        raise HTTPException(status_code=413, detail=f"Input exceeds {MAX_FILE_SIZE // (1024 * 1024)} MB limit")
    if not data:
        raise HTTPException(status_code=400, detail="Empty upload")
    return data, suffix


def _json_safe(value: Any) -> Any:
    if isinstance(value, dict):
        return {str(k): _json_safe(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_safe(v) for v in value]
    if hasattr(value, "tolist"):
        return _json_safe(value.tolist())
    if isinstance(value, Path):
        return str(value)
    try:
        json.dumps(value)
        return value
    except TypeError:
        return str(value)


def _result_json(result: Any) -> dict[str, Any]:
    payload = getattr(result, "json", None)
    if callable(payload):
        payload = payload()
    if payload is None:
        payload = getattr(result, "res", None)
    if payload is None:
        payload = {"result": str(result)}
    return _json_safe(payload)


def _extract_text(payloads: list[dict[str, Any]]) -> str:
    lines: list[str] = []

    def walk(value: Any) -> None:
        if isinstance(value, dict):
            for key, child in value.items():
                if key in {"rec_texts", "rec_text", "text", "content"}:
                    if isinstance(child, list):
                        lines.extend(str(item) for item in child if str(item).strip())
                    elif str(child).strip():
                        lines.append(str(child))
                else:
                    walk(child)
        elif isinstance(value, list):
            for child in value:
                walk(child)

    for payload in payloads:
        walk(payload)
    return "\n".join(dict.fromkeys(lines))


def _tesseract_enabled() -> bool:
    return os.getenv("OCR_TESSERACT_FALLBACK", "1").strip().lower() not in {"0", "false", "no", "off"}


def _tesseract_text(input_path: Path, suffix: str) -> tuple[str, list[dict[str, Any]]]:
    import pytesseract
    from PIL import Image

    lang = os.getenv("TESSERACT_LANG", "vie+eng")
    images: list[Any] = []
    if suffix == ".pdf":
        import pypdfium2 as pdfium

        pdf = pdfium.PdfDocument(str(input_path))
        try:
            for index in range(len(pdf)):
                page = pdf[index]
                bitmap = page.render(scale=2.2)
                images.append(bitmap.to_pil())
        finally:
            pdf.close()
    else:
        images.append(Image.open(input_path).convert("RGB"))

    pages: list[dict[str, Any]] = []
    texts: list[str] = []
    try:
        for index, image in enumerate(images, start=1):
            text = pytesseract.image_to_string(image, lang=lang).strip()
            texts.append(text)
            pages.append({"page": index, "text": text, "engine": "tesseract", "language": lang})
    finally:
        for image in images:
            closer = getattr(image, "close", None)
            if callable(closer):
                closer()
    return "\n\n".join(text for text in texts if text), pages


def _safe_sheet_name(base: str, used: set[str]) -> str:
    cleaned = "".join("_" if ch in "[]:*?/\\" else ch for ch in base).strip() or "Table"
    cleaned = cleaned[:31]
    candidate = cleaned
    index = 2
    while candidate in used:
        suffix = f"_{index}"
        candidate = f"{cleaned[: 31 - len(suffix)]}{suffix}"
        index += 1
    used.add(candidate)
    return candidate


def _copy_style(source: Any, target: Any) -> None:
    """Copy a style across workbooks without reusing foreign style-table IDs."""
    if not getattr(source, "has_style", False):
        return
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.number_format = source.number_format
    target.protection = copy(source.protection)
    if hasattr(source, "quotePrefix") and hasattr(target, "quotePrefix"):
        target.quotePrefix = source.quotePrefix
    if hasattr(source, "pivotButton") and hasattr(target, "pivotButton"):
        target.pivotButton = source.pivotButton


def _copy_sheet(source: Any, target: Any, row_offset: int = 0) -> None:
    from openpyxl.utils.cell import range_boundaries

    for row in source.iter_rows():
        for cell in row:
            dst = target.cell(row=cell.row + row_offset, column=cell.column)
            dst.value = cell.value
            _copy_style(cell, dst)
            if cell.hyperlink:
                dst._hyperlink = copy(cell.hyperlink)
            if cell.comment:
                dst.comment = copy(cell.comment)

    for merged_range in source.merged_cells.ranges:
        min_col, min_row, max_col, max_row = range_boundaries(str(merged_range))
        target.merge_cells(
            start_row=min_row + row_offset,
            start_column=min_col,
            end_row=max_row + row_offset,
            end_column=max_col,
        )

    for key, dimension in source.column_dimensions.items():
        target_dimension = target.column_dimensions[key]
        if dimension.width is not None:
            target_dimension.width = max(target_dimension.width or 0, dimension.width)
        target_dimension.hidden = bool(target_dimension.hidden or dimension.hidden)
        target_dimension.bestFit = bool(target_dimension.bestFit or dimension.bestFit)
        target_dimension.outlineLevel = max(target_dimension.outlineLevel or 0, dimension.outlineLevel or 0)
        target_dimension.collapsed = bool(target_dimension.collapsed or dimension.collapsed)
        target_dimension.min = dimension.min
        target_dimension.max = dimension.max
        _copy_style(dimension, target_dimension)

    for key, dimension in source.row_dimensions.items():
        target_dimension = target.row_dimensions[key + row_offset]
        target_dimension.height = dimension.height
        target_dimension.hidden = dimension.hidden
        target_dimension.outlineLevel = dimension.outlineLevel
        target_dimension.collapsed = dimension.collapsed
        if hasattr(dimension, "thickTop"):
            target_dimension.thickTop = dimension.thickTop
        if hasattr(dimension, "thickBot"):
            target_dimension.thickBot = dimension.thickBot
        _copy_style(dimension, target_dimension)

    if row_offset == 0:
        target.freeze_panes = source.freeze_panes
        target.sheet_format = copy(source.sheet_format)
        target.sheet_properties = copy(source.sheet_properties)
        target.page_margins = copy(source.page_margins)
        target.page_setup = copy(source.page_setup)
        target.print_options = copy(source.print_options)


def _source_page_sizes(input_path: Path, suffix: str) -> dict[int, tuple[float, float]]:
    """Return source page/image dimensions for print-layout preservation.

    Scanned PDFs often carry custom page dimensions rather than a standard paper
    code. Only the aspect ratio is needed: the XLSX print sheet is normalized to
    a practical 14-inch long edge while preserving the source page orientation.
    """

    try:
        if suffix == ".pdf":
            import pypdfium2 as pdfium

            pdf = pdfium.PdfDocument(str(input_path))
            try:
                return {
                    index + 1: (float(pdf[index].get_width()), float(pdf[index].get_height()))
                    for index in range(len(pdf))
                }
            finally:
                pdf.close()

        from PIL import Image

        with Image.open(input_path) as image:
            return {1: (float(image.width), float(image.height))}
    except Exception:
        return {}


def _configure_page_print_layout(target: Any, page_size: tuple[float, float] | None) -> None:
    """Keep one recognized source page printable as one aspect-correct sheet."""

    target.page_setup.scale = None
    target.page_setup.fitToWidth = 1
    target.page_setup.fitToHeight = 1
    target.sheet_properties.pageSetUpPr.fitToPage = True
    target.page_margins.left = 0.2
    target.page_margins.right = 0.2
    target.page_margins.top = 0.3
    target.page_margins.bottom = 0.3

    if page_size is not None:
        width, height = page_size
        if width > 0 and height > 0:
            long_edge = max(width, height)
            paper_width = PRINT_LONG_EDGE_INCHES * width / long_edge
            paper_height = PRINT_LONG_EDGE_INCHES * height / long_edge
            target.page_setup.paperSize = None
            target.page_setup.paperWidth = f"{paper_width:.4f}in"
            target.page_setup.paperHeight = f"{paper_height:.4f}in"
            target.page_setup.orientation = "landscape" if width > height else "portrait"

    target.print_area = target.calculate_dimension()


def _merge_xlsx_files(
    generated: list[tuple[int, Path]],
    page_sizes: dict[int, tuple[float, float]] | None = None,
) -> bytes:
    """Merge PP-Structure XLSX outputs into one worksheet per source page.

    PP-Structure may emit multiple workbooks/tables for one page. Keeping those
    tables on separate worksheets makes a single scanned page print as multiple
    pages. SolYan instead stacks tables from the same source page vertically,
    preserves their styles/merged cells, and applies one-page print scaling.
    """

    from openpyxl import Workbook, load_workbook

    output = Workbook()
    output.remove(output.active)
    used_names: set[str] = set()
    page_sizes = page_sizes or {}

    grouped: dict[int, list[Path]] = {}
    for page_index, path in generated:
        grouped.setdefault(page_index, []).append(path)

    for page_index, paths in grouped.items():
        target_sheet = output.create_sheet(_safe_sheet_name(f"Page {page_index}", used_names))
        row_offset = 0

        for path in paths:
            source_book = load_workbook(path, data_only=False)
            try:
                for source_sheet in source_book.worksheets:
                    _copy_sheet(source_sheet, target_sheet, row_offset=row_offset)
                    row_offset += max(source_sheet.max_row, 1) + TABLE_GAP_ROWS
            finally:
                source_book.close()

        _configure_page_print_layout(target_sheet, page_sizes.get(page_index))

    if not output.worksheets:
        raise ValueError("PP-StructureV3 produced no table worksheets")

    stream = io.BytesIO()
    output.save(stream)
    output.close()
    return stream.getvalue()


def _structure_results(input_path: Path) -> list[Any]:
    try:
        return list(_load_structure().predict(input=str(input_path)))
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Structured OCR failed: {exc}") from exc


@app.get("/health")
def health() -> dict[str, object]:
    return {
        "status": "ok",
        "engine": "PaddleOCR",
        "models_loaded": {"ocr": _ocr_pipeline is not None, "structure": _structure_pipeline is not None},
        "fallback": {
            "enabled": _tesseract_enabled(),
            "engine": "Tesseract",
            "binary_available": shutil.which("tesseract") is not None,
            "language": os.getenv("TESSERACT_LANG", "vie+eng"),
        },
        "exports": {"xlsx": True},
    }


@app.post("/ocr/text")
def ocr_text(file: UploadFile = File(...)) -> dict[str, Any]:
    data, suffix = _read_upload(file)
    with tempfile.TemporaryDirectory(prefix="solyan-ocr-") as temp_dir:
        input_path = Path(temp_dir) / f"input{suffix}"
        input_path.write_bytes(data)
        try:
            results = list(_load_ocr().predict(str(input_path)))
            pages = [_result_json(result) for result in results]
            return {"mode": "text", "engine": "PaddleOCR", "pages": pages, "text": _extract_text(pages), "fallback_used": False}
        except Exception as paddle_exc:
            if not _tesseract_enabled():
                raise HTTPException(status_code=422, detail=f"OCR failed: {paddle_exc}") from paddle_exc
            try:
                text, pages = _tesseract_text(input_path, suffix)
                return {
                    "mode": "text",
                    "engine": "Tesseract",
                    "pages": pages,
                    "text": text,
                    "fallback_used": True,
                    "primary_error": str(paddle_exc)[:1200],
                }
            except Exception as fallback_exc:
                raise HTTPException(
                    status_code=422,
                    detail=f"PaddleOCR failed ({paddle_exc}); Tesseract fallback failed ({fallback_exc})",
                ) from fallback_exc


@app.post("/ocr/structured")
def ocr_structured(file: UploadFile = File(...)) -> dict[str, Any]:
    data, suffix = _read_upload(file)
    with tempfile.TemporaryDirectory(prefix="solyan-structure-") as temp_dir:
        root = Path(temp_dir)
        input_path = root / f"input{suffix}"
        output_dir = root / "output"
        output_dir.mkdir()
        input_path.write_bytes(data)
        results = _structure_results(input_path)
        payloads = [_result_json(result) for result in results]
        try:
            for result in results:
                saver = getattr(result, "save_to_markdown", None)
                if callable(saver):
                    saver(save_path=str(output_dir))
        except Exception as exc:
            raise HTTPException(status_code=422, detail=f"Structured OCR export failed: {exc}") from exc

        markdown_parts: list[str] = []
        for path in sorted(output_dir.rglob("*.md")):
            try:
                markdown_parts.append(path.read_text(encoding="utf-8"))
            except UnicodeDecodeError:
                continue

        return {
            "mode": "structured",
            "engine": "PP-StructureV3",
            "pages": payloads,
            "markdown": "\n\n---\n\n".join(part for part in markdown_parts if part.strip()),
        }


@app.post("/ocr/xlsx")
def ocr_xlsx(file: UploadFile = File(...)) -> Response:
    """Recognize scanned tables and return one XLSX workbook.

    PP-StructureV3 may emit one or more workbooks per page/table. SolYan merges
    those outputs into one worksheet per source page, stacking multiple detected
    tables on the same page while preserving source-page print orientation/aspect.
    """
    data, suffix = _read_upload(file)
    with tempfile.TemporaryDirectory(prefix="solyan-xlsx-") as temp_dir:
        root = Path(temp_dir)
        input_path = root / f"input{suffix}"
        output_dir = root / "tables"
        output_dir.mkdir()
        input_path.write_bytes(data)

        results = _structure_results(input_path)
        generated: list[tuple[int, Path]] = []
        for page_index, result in enumerate(results, start=1):
            saver = getattr(result, "save_to_xlsx", None)
            if not callable(saver):
                continue
            page_dir = output_dir / f"page-{page_index:03d}"
            page_dir.mkdir()
            try:
                saver(save_path=str(page_dir))
            except Exception as exc:
                raise HTTPException(
                    status_code=422,
                    detail=f"PP-StructureV3 XLSX export failed on page {page_index}: {exc}",
                ) from exc
            generated.extend((page_index, path) for path in sorted(page_dir.rglob("*.xlsx")))

        if not generated:
            raise HTTPException(status_code=422, detail="PP-StructureV3 found no exportable tables")

        try:
            workbook = _merge_xlsx_files(generated, page_sizes=_source_page_sizes(input_path, suffix))
        except Exception as exc:
            traceback.print_exc()
            generated_names = ", ".join(f"page={page_index}:{path.name}" for page_index, path in generated)
            raise HTTPException(
                status_code=422,
                detail=(
                    f"Unable to assemble OCR XLSX workbook: {type(exc).__name__}: {exc}; "
                    f"generated=[{generated_names}]"
                ),
            ) from exc

        return Response(
            content=workbook,
            media_type=XLSX_MEDIA_TYPE,
            headers={
                "Content-Disposition": 'attachment; filename="solyan-ocr-tables.xlsx"',
                "X-SolYan-OCR-Engine": "PP-StructureV3",
                "X-SolYan-OCR-Pages": str(len(results)),
                "X-SolYan-OCR-Table-Workbooks": str(len(generated)),
            },
        )
